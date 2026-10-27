"""FIX B52: báo cáo ngày buổi học, không sửa database hoặc đọc credential."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import zipfile

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.75)
s.left_margin = s.right_margin = Inches(.8)
for name, size in [('Normal', 11), ('Title', 22), ('Heading 1', 16), ('Heading 2', 12)]:
    st = doc.styles[name]
    st.font.name, st.font.size = 'Arial', Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.08
doc.core_properties.title = 'Báo cáo sửa lỗi B52 ngày buổi học không có múi giờ'
doc.core_properties.author = 'SportCenter'
def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t, 2)
def code(t):
    q=doc.add_paragraph()
    r=q.add_run(t)
    r.font.name, r.font.size = 'Consolas', Pt(9)

doc.add_heading('Báo cáo sửa lỗi B52 ngày buổi học', 0)
p('Ngày 27 tháng 10 năm 2026')
p('Đã sửa ngày buổi học trong response lịch lớp và điểm danh để có múi giờ UTC rõ ràng. Giữ kiểu DateTime của DTO và dữ liệu DateOnly trong DB; không đổi giờ học, không đổi quyền hoặc logic đặt lớp. Chỉ sửa B52, các nghiệp vụ B45 đến B51 trong file đính kèm chưa thay đổi ở đợt này.')
h('Sai ở đâu và vì sao hiển thị lùi một ngày')
p('ClassSession.Date lưu DateOnly, tức ngày lịch không có giờ hay múi giờ. Khi tạo DTO, code cũ gọi ToDateTime(TimeOnly.MinValue) nên DateTime.Kind là Unspecified. JSON thành 2026-10-15T00:00:00, không có hậu tố Z hoặc offset.')
p('Nếu trình duyệt tại Việt Nam hiểu chuỗi này là 00:00 ngày 15/10 theo UTC+7, thời điểm tương ứng là 17:00 ngày 14/10 UTC. Khi FE tiếp tục định dạng ngày theo UTC thì có thể hiện 14/10. Lỗi nằm ở hợp đồng serialize ngày thiếu thông tin, không phải DB tự thay ngày 15 thành 14.')
h('Cách sửa đã chọn')
code('// CODE CŨ\nDate = cs.Date.ToDateTime(TimeOnly.MinValue);\n\n// FIX B52\nDate = DateTime.SpecifyKind(\n    cs.Date.ToDateTime(TimeOnly.MinValue), DateTimeKind.Utc);')
p('Response mới là 2026-10-15T00:00:00Z. Thêm Kind UTC giúp FE không phải suy đoán theo giờ máy. Chọn cách này thay vì đổi response sang DateOnly dạng yyyy-MM-dd để giữ tương thích với DTO DateTime và các client đang đọc chuỗi ngày giờ.')
p('SpecifyKind gắn quy ước UTC cho giá trị ngày, không chuyển 00:00 Việt Nam thành UTC và không cộng/trừ bảy giờ. Đây là biểu diễn ngày lịch qua DateTime, không khẳng định buổi học thực tế bắt đầu lúc 00:00 UTC. Giờ học thật vẫn lấy StartTime và EndTime theo nghiệp vụ hiện có.')
h('Code cũ được giữ thế nào')
p('Tại từng chỗ sửa có comment CODE CŨ và FIX B52 để đối chiếu. Không xóa entity DateOnly, không chỉnh AppDbContext, không sửa dữ liệu lịch sử và không thay global JSON serializer.')

doc.add_page_break()
doc.add_heading('Phạm vi kiểm thử và triển khai', 1)
h('Sáu chỗ ánh xạ đã sửa')
p('ClassSessionService.GetAllAsync: trường Date của danh sách tổng. GetByClassIdAsync: Date của các buổi theo lớp. GetByIdAsync: Date của chi tiết; response sau tạo buổi dùng lại hàm này nên cũng nhận định dạng mới.')
p('AttendanceService.GetBySessionIdAsync: SessionDate khi đọc điểm danh theo buổi. GetByMemberIdAsync: SessionDate của lịch sử Member, bao gồm alias /api/members/me/attendances. RecordAttendanceAsync: SessionDate trong response sau ghi điểm danh. Giá trị fallback DateTime.MinValue ở hai projection đọc cũng được gắn UTC.')
h('Kết quả kiểm tra')
p('Build bộ hồi quy thành công. Tổng 201 kiểm tra qua, gồm 9 kiểm tra B52 mới: sáu response HTTP trả chính xác 2026-10-15T00:00:00Z với Kind UTC; thao tác ghi fixture điểm danh thành công; DateOnly trong DB test vẫn là 15/10; provider Npgsql biên dịch được projection ngày thành truy vấn SQL mà không kết nối database thật.')
p('Test đi qua Controller, Service và JSON serializer bằng database InMemory riêng. Các kiểm tra cũ về quyền, thanh toán, upload và nghiệp vụ khác vẫn chạy. Không gọi database Render hoặc giả định kết quả hiển thị trên FE production đã được xác minh. Kiểm tra ToQueryString xác nhận truy vấn biên dịch, không thay thế smoke test thực thi PostgreSQL.')
code('dotnet build tests/SportCenter.Regression/SportCenter.Regression.csproj\n  --no-restore -p:BuildInParallel=false -m:1\n  -p:UseSharedCompilation=false -nodeReuse:false\n\ndotnet tests/SportCenter.Regression/bin/Debug/net8.0/SportCenter.Regression.dll\n\nGhép các dòng của lệnh build thành một dòng khi chạy PowerShell.')
h('Có cần migration hoặc Update Database không')
p('Không. B52 chỉ sửa cách tạo response, không thay schema hay giá trị trong DB. Sau khi review code, commit/push và deploy lại API theo quy trình của nhóm. Đợt sửa này không tự commit, push, deploy hoặc chạy database update.')
h('FE cần kiểm tra sau deploy')
p('Trong Network, mở API class-sessions và attendances, xác nhận Date/SessionDate có Z. Với dữ liệu ngày 15/10, định dạng ngày theo UTC phải giữ 15/10. Kiểm tra trang lớp, lịch học, điểm danh và trang HLV đang tiêu thụ các response này. Có thể giữ fallback đọc chuỗi không múi giờ của FE để tương thích dữ liệu cũ; chuỗi mới có Z phải được parse như UTC bình thường.')
p('Không tự chuyển StartTime/EndTime từ giờ Việt Nam sang UTC chỉ vì trường ngày có Z. Nếu sau này nhóm đổi toàn bộ hợp đồng sang DateOnly, cần thống nhất cả FE và API thay vì chỉ sửa một màn hình.')
p('File nguồn thay đổi: SWP_SportCenter.Service/ClassSession/ClassSessionService.cs; SWP_SportCenter.Service/Attendance/AttendanceService.cs; tests/SportCenter.Regression/Program.cs.')

out=Path(__file__).with_name('Bao_cao_sua_loi_B52_ngay_buoi_hoc.docx')
doc.save(out)
with zipfile.ZipFile(out) as z: assert z.testzip() is None
assert len(Document(out).paragraphs) > 25
print(out)
